"""Import the FM asset register into the SOP app as sop/register.js.

Reads data/Master_Asset_Register_for_FM_Sep_2026.xlsx (sheets "Asset Master
List", "House Equip", "FR Furniture" and "dorm parents"; "Sheet1" lists
residents' names and is skipped), cleans the rows, removes duplicate records,
merges identical unit rows, derives each asset's trade system and writes:

    sop/register.js                        window.SOP_REGISTER (sop/REGISTER_SCHEMA.md)
    dist/Asset_Register_Dedup_Report.xlsx  what was removed, merged or may overlap

Duplicate rules:
  A. Same meaningful serial (Identification # S.No) or same single-unit Asset
     Tag #, within or across sheets, AND same building, location, room and
     description (model may be blank on one side): a true double entry, so the
     later row is dropped. Rows that share an ID but differ otherwise (or carry
     their own different serial/tag) are kept, each with a `note` naming the other.
  B. Rows with no serial/unique tag that are identical in building, location,
     room, category, description, model, tag (and the PM fields) are units of
     one item: merged into one entry, qty summed.
  C. FR Furniture / dorm parents vs Asset Master List: removed only on an ID
     match (rule A); same-residence look-alikes go to "Possible overlaps".

Occupant names never leave the workbook: House Equip "Room No." cells that hold
a name instead of a room code are dropped, unnamed columns are never read, and
the run fails if any name from Sheet1 or those columns reaches an output.

    python3 tools/import_register.py
"""
import collections
import datetime
import json
import pathlib
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "Master_Asset_Register_for_FM_Sep_2026.xlsx"
OUT_JS = ROOT / "sop" / "register.js"
OUT_REPORT = ROOT / "dist" / "Asset_Register_Dedup_Report.xlsx"

MASTER, HOUSE, FR, DORM = "Asset Master List", "House Equip", "FR Furniture", "dorm parents"
SHEETS = [MASTER, HOUSE, FR, DORM]          # also the "first row wins" order across sheets
EXCLUDED = "Sheet1"                          # residents' names and addresses: personal data

FREQ = {  # code: (label, days)
    "M": ("Monthly", 30), "Bi": ("Bimonthly", 60), "Q": ("Quarterly", 90),
    "H": ("Half-yearly", 180), "Y": ("Yearly", 365),
}
CRITICALITY = {"C": "Critical", "NC": "Non-critical"}
FURNITURE_CATS = {"Furniture Fixture", "Furniture"}

# ---------------------------------------------------------------- system rules
# First match wins. Each rule: (system, regex on "description category", only
# for these categories or None). Keep them readable; tune with the printout.
F = re.I
FURN_WORDS = (r"\b(chairs?|stools?|tables?|desks?|beds?|sofas?|cupboards?|ward ?robes?|war drop|almirah|"
              r"racks?|shel(f|ves|ving)|bookshel|lockers?|loacable|benche?s?|cabinets?|mirror|drawers?|dwrs?|"
              r"sideboard|tv unit|tv table|podium|lectern|stand|whiteboard|white board|notice board|pin board|"
              r"board|partition|troll(e)?y|mattres|matters|curtains?|blinds?|carpet|dustbin|bins?|easel|"
              r"tablet|cantilever|workstation|counter|step by step|i-2)\b")
RULES = [
    ("Solar", r"\bsolar\b", None),
    ("Fire safety", r"\bfire\b|extingu|hydrant|sprinkler|smoke|alarm panel|\b(abc|co2) type\b|^water \(\d+ ?kg\)", None),
    ("Lifts", r"\blifts?\b|elevator", None),
    ("Pool", r"\bpool\b|sand filter", None),
    ("Water treatment", r"soft[e]?ner|\bwtp\b|\bstp\b|ro plant|chlorin|\bblower|\bmgf\b|\bacf\b|filter feed", None),
    # Domestic appliances in residences (House Equip) and in the master list.
    ("Appliances", r"chimn|gas stove|induction|micro ?wave|refrigerat|fridge|washing machine|purifier|kent ro|"
                   r"\bro\b|water cooler|dispens|air cooler|\btv\b|television|\biron\b|tumble dryer|"
                   r"hydro machine|boiler machine", {"House equipment", "Equipment"}),
    ("Appliances", r"w(a)?ter dispens", None),
    ("Kitchen", r"\bsink\b", None),
    ("Furniture", FURN_WORDS, FURNITURE_CATS),
    ("Kitchen", r"kitchen|oven|fryer|\bhood\b|dish ?wash|co+ld room|ice cube|grinder|atta machine|peeler|"
                r"gas range|dosa|t[o]?aster|toster|braed|bread|vegetable cutter|benberry|bain marie|"
                r"masala|mixer|bakery|\bss work table", None),
    ("HVAC", r"\bac\b|air condition|split|cassette|chiller|\bahu\b|\bfcu\b|\bvrf\b|cooling tower|exhaust|"
             r"ventilat|ductable|\btfa\b|heat pump|fab motor", None),
    ("Electrical", r"\bdg\b|generator|transformer|panel|\bups\b|breaker|brekar|inverter|capacitor|kvar|"
                   r"light|meter|electric|\bdb'?s?\b|feeder pillar|\bamf\b|switch|cable", None),
    ("Plumbing", r"pump|motor|bore ?well|tank|geyser|heater|plumb|flush|\btaps?\b|valve", None),
    ("IT & AV", r"projector|computer|laptop|printer|speaker|cctv|camera|network|\bpa system|amplifier|"
                r"monitor|screen|mic\b", None),
    ("Lab equipment", r"\blab\b|laboratory|fume|microscope|autoclave|vac+um table", None),
    ("Sports", r"\bgym|treadmill|sports?\b|squash|basketball|football|cricket|badminton", None),
    ("Building fabric", r"\bwalls?\b|\broof|\bdoors?\b|windows?|floor(ing)?\b", None),
    ("Building fabric", r"\bBuilding$", {"Building"}),
    ("Furniture", r".", FURNITURE_CATS),
    ("Appliances", r".", {"House equipment"}),
]
RULES = [(s, re.compile(p, F), cats) for s, p, cats in RULES]


def system_for(name, category):
    text = f"{name} {category}"
    for system, rx, cats in RULES:
        if (cats is None or category in cats) and rx.search(text):
            return system
    return "General"


# ---------------------------------------------------------------- cleaning
PLACEHOLDER = {"", "-", "--", "NA", "N/A", "NIL", "NONE", "NULL"}
NO_ID = PLACEHOLDER | {"0", "NO", "DAMAGED", "SHIFTED", "NOTAVAILABLE"}


def text(v):
    """Trimmed, single-spaced string; placeholders become ''."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return "" if s.upper() in PLACEHOLDER else s


def number(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else round(float(v), 3)
    m = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*", str(v))
    return number(float(m.group(1))) if m else None


def quantity(v):
    """Units a row covers. Blank/0 count as 1; '2+1' is 3."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        n = int(round(v))
    else:
        n = sum(int(x) for x in re.findall(r"\d+", str(v or "")))
    return n if n > 0 else 1


def freq_code(v):
    s = text(v).upper()
    if not s or s in ("N", "NO"):
        return None
    if s.startswith("BI"):
        return "Bi"
    for code in ("M", "Q", "H", "Y"):
        if s.startswith(code):
            return code
    return None


def yes(v):
    return text(v).upper() in ("Y", "YES")


def key(v):
    """Comparison form of a serial/tag: no spaces, upper case."""
    return re.sub(r"\s+", "", text(v)).upper()


def meaningful(k):
    return (k not in NO_ID and len(k) >= 4
            and not re.fullmatch(r"\d{1,3}([,/&+-]\d{1,3})*", k))   # 12, 007,008: sequence numbers


def range_tag(k):
    """AKF/ACADMY/FUR/2-5 style tags cover several units, so they are not unique."""
    return bool(re.search(r"\d+-\d+$", k) or "," in k)


def alnum(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def has_digit(s):
    return bool(re.search(r"\d", s))


# ---------------------------------------------------------------- reading
def header_row(rows):
    for i, r in enumerate(rows):
        cells = [text(c) for c in r]
        if "Description" in cells and "Building Name" in cells:
            return i, cells
    raise SystemExit("header row not found")


def read_sheet(ws, private_names):
    rows = list(ws.iter_rows(values_only=True))
    hi, hdr = header_row(rows)
    cols = {h: i for i, h in reversed(list(enumerate(hdr))) if h}   # first column of each named header
    out = []
    for n, r in enumerate(rows[hi + 1:], hi + 2):
        get = lambda h, r=r: r[cols[h]] if h in cols and cols[h] < len(r) else None
        if not text(get("Description")):
            continue
        # Columns without a header (dorm parents col K, House Equip "Name of Occupant") hold names.
        for i, h in enumerate(hdr):
            if not h and i < len(r) and isinstance(r[i], str) and text(r[i]) and not has_digit(r[i]):
                private_names.add(text(r[i]))
        out.append((n, get))
    return out


def row_record(sheet, n, get, private_names):
    rec = {"sheet": sheet, "rows": [n], "name": text(get("Description")), "building": text(get("Building Name")),
           "room": text(get("Room No.")), "serial": text(get("Identification # S.No")),
           "units": quantity(get("Quantity"))}
    rec["units0"] = rec["units"]
    if sheet in (FR, DORM):
        rec["category"] = "Furniture"
        rec["location"] = text(get("Area"))
        rec["room"] = rec["room"].strip("()")
        rec["model"] = " · ".join(x for x in (text(get("Code")), text(get("Size"))) if x)
        rec["tag"] = ""
    else:
        rec["category"] = "House equipment" if sheet == HOUSE else text(get("Category"))
        rec["location"] = text(get("Location"))
        rec["model"] = text(get("Model"))
        rec["tag"] = text(get("Asset Tag #"))
        year = number(get("Year of Purchase"))
        rec["year"] = year if year and year > 1900 else None
        rec["life"] = number(get("life cycle (in years)")) or None
        pct = number(get("% Operational life"))
        rec["lifePct"] = round(pct, 3) if isinstance(pct, (int, float)) and rec["year"] else None
        rec["criticality"] = CRITICALITY.get(text(get("Critical/ Non Critical")).upper(), "")
        amc, ppm = freq_code(get("AMC frequency")), freq_code(get("PPM frequency"))
        rec["amcCode"] = amc if yes(get("AMC")) else None
        rec["ppmCode"] = ppm if yes(get("Covered under PPM")) else None
        rec["amc"] = (FREQ[amc][0] if amc else "Yes") if yes(get("AMC")) else ""
        rec["ppm"] = (FREQ[ppm][0] if ppm else "Yes") if yes(get("Covered under PPM")) else ""
        if sheet == HOUSE and rec["room"] and not has_digit(rec["room"]):
            private_names.add(rec["room"])      # occupant's name typed into Room No.
            rec["room"] = ""
    if rec["serial"].upper() in NO_ID - {"DAMAGED", "SHIFTED"}:
        rec["serial"] = ""
    sk, tk = key(rec["serial"]), key(rec["tag"])
    # Serials compare without punctuation (VVA0119 = VVA-0119); tags keep it so ranges stay detectable.
    rec["keys"] = ([("serial", re.sub(r"[^A-Z0-9]", "", sk))] if meaningful(sk) else []) + \
                  ([("tag", tk)] if meaningful(tk) and not range_tag(tk) else [])
    return rec


# ---------------------------------------------------------------- helpers
def src(rec):
    nums = sorted(rec["rows"])
    parts, start = [], nums[0]
    for a, b in zip(nums, nums[1:] + [None]):
        if b != a + 1:
            parts.append(f"r{start}" if start == a else f"r{start}-r{a}")
            start = b
    return f"{rec['sheet']} " + ", ".join(parts)


def place(rec):
    return " / ".join(x for x in (rec["building"], rec["location"], rec["room"]) if x)


FILLABLE = ["model", "serial", "tag", "year", "life", "lifePct", "criticality", "amc", "ppm", "amcCode", "ppmCode",
            "location", "room"]


APPLIANCE_TYPES = [("Refrigerator", r"refrigerat|fridge"), ("Microwave", r"micro ?wave"),
                   ("Induction cooker", r"induction"), ("Kitchen chimney", r"chimn"), ("Gas stove", r"gas stove"),
                   ("Washing machine", r"washing machine"), ("Water purifier", r"purifier|kent ro"),
                   ("Air conditioner", r"air condition|split ac")]


def residence(s):
    """FR3B, DFR1, SR11A ... from free-typed building names (' SR-1', 'FR-3B Left', 'SR9,SR10 B/A')."""
    m = re.search(r"(DFR|FR|SR)(\d+)([A-E]?)", re.sub(r"[\s-]", "", s.upper()))
    return (m.group(1), m.group(2), m.group(3)) if m else None


def same_residence(a, b):
    return a and b and a[:2] == b[:2] and (a[2] == b[2] or not a[2] or not b[2])


# ---------------------------------------------------------------- main
def main():
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    private_names = set()
    # Names in the excluded sheet are only collected to prove they never reach an output.
    if EXCLUDED in wb.sheetnames:
        for r in wb[EXCLUDED].iter_rows(values_only=True):
            for c in r:
                if isinstance(c, str) and text(c) and not has_digit(c) and text(c).lower() not in (
                        "name", "address", "sl no", "sl no.", "fridge", "washing machine", "micro wave"):
                    private_names.add(text(c))
    records, rows_read = [], collections.Counter()
    for sheet in SHEETS:
        for n, get in read_sheet(wb[sheet], private_names):
            records.append(row_record(sheet, n, get, private_names))
            rows_read[sheet] += 1
    total_units = sum(r["units0"] for r in records)

    # --- A: same serial / unique tag -> one physical asset
    parent = list(range(len(records)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    first = {}
    for i, rec in enumerate(records):
        for k in rec["keys"]:
            if k[1] in first:
                a, b = find(first[k[1]]), find(i)
                parent[max(a, b)] = min(a, b)
            else:
                first[k[1]] = i
    groups = collections.defaultdict(list)
    for i in range(len(records)):
        groups[find(i)].append(i)
    ids = lambda r: dict(r["keys"])

    def why_separate(a, b):
        """Reasons two rows sharing an ID are separate units; [] means a true double entry."""
        why = []
        for f, label in (("building", "building"), ("location", "location/area"), ("room", "room"),
                         ("name", "description")):
            if alnum(a[f]) != alnum(b[f]):
                why.append(f"{label} '{a[f]}' vs '{b[f]}'")
        if a["model"] and b["model"] and alnum(a["model"]) != alnum(b["model"]):
            why.append(f"model '{a['model']}' vs '{b['model']}'")
        for kind in ("serial", "tag"):   # a second, different ID means two physical units
            if kind in ids(a) and kind in ids(b) and ids(a)[kind] != ids(b)[kind]:
                why.append(f"different {kind}s '{a[kind]}' vs '{b[kind]}'")
        return why

    def shared(a, b):
        return [(kind, b[kind]) for kind, k in b["keys"] if k in set(ids(a).values())]

    removed, collisions = [], []      # true double entries; kept rows that share an ID
    for root, members in groups.items():
        if len(members) < 2:
            continue
        kept = []
        for i in members:
            rec = records[i]
            twin = next((k for k in kept if shared(k, rec) and not why_separate(k, rec)), None)
            if twin is None:
                kept.append(rec)
                continue
            for f in FILLABLE:
                if not twin.get(f) and rec.get(f):
                    twin[f] = rec[f]
            rec["removed"] = True
            removed.append({"rec": rec, "kept": twin,
                            "match": ", ".join(f"{k} {v}" for k, v in shared(twin, rec)),
                            "note": ("Cross-sheet. " if rec["sheet"] != twin["sheet"] else "") +
                                    "Same asset recorded twice (same ID, building, location, room and description)."})
        for x, a in enumerate(kept):
            for b in kept[x + 1:]:
                if shared(a, b):
                    collisions.append((a, b, shared(a, b), why_separate(a, b)))

    # --- B: identical rows without an ID -> one entry, qty summed
    entries, merged_index = [], {}
    for rec in records:
        if rec.get("removed"):
            continue
        if rec["keys"]:
            entries.append(rec)
            continue
        serial = "" if re.fullmatch(r"\d{1,3}", key(rec["serial"])) else key(rec["serial"])
        k = (rec["sheet"], rec["building"].lower(), rec["location"].lower(),
             rec["room"].lower(), rec["category"], rec["name"].lower(), rec["model"].lower(), key(rec["tag"]),
             serial, rec.get("year"), rec.get("life"), rec.get("criticality"), rec.get("amc"), rec.get("ppm"))
        if k in merged_index:
            m = merged_index[k]
            m["rows"] += rec["rows"]
            m["units"] += rec["units"]
            if m["serial"] != rec["serial"]:
                m["serial"] = ""
        else:
            merged_index[k] = rec
            entries.append(rec)
    merged = [e for e in entries if len(e["rows"]) > 1]

    # --- IDs and systems
    counters = collections.Counter()
    for e in entries:
        if e["sheet"] == MASTER:
            prefix, width = {"Equipment": ("EQ", 4), "Building": ("BL", 4)}.get(e["category"], ("FF", 5))
        else:
            prefix, width = {HOUSE: "HE", FR: "FR", DORM: "DP"}[e["sheet"]], 4
        counters[prefix] += 1
        e["id"] = f"{prefix}-{counters[prefix]:0{width}d}"
        e["system"] = system_for(e["name"], e["category"])
        code = e.get("ppmCode") or e.get("amcCode")
        e["freq"] = FREQ[code][1] if code else (365 if e["category"] in FURNITURE_CATS else 180)

    # Kept collisions: both assets carry a data-quality note naming the other one.
    for a, b, sh, _ in collisions:
        for me, other in ((a, b), (b, a)):
            what = " and ".join(f"{kind.capitalize() if not i else kind} {v}" for i, (kind, v) in enumerate(sh))
            where = " ".join(x for x in (other["building"], other["room"]) if x)
            me.setdefault("notes", []).append(
                f"{what} is also recorded on {other['id']} ({where} {other['name']}). Verify the "
                f"{'serial' if sh[0][0] == 'serial' else 'asset tag'}.")
    # --- C: furniture lists vs master furniture in the same residence
    master_furn = [e for e in entries if e["sheet"] == MASTER and e["category"] == "Furniture Fixture"]
    for e in master_furn:
        e["res"] = residence(e["building"]) or residence(e["building"] + " " + e["location"])
        e["desc"] = alnum(e["name"])
    overlaps = []
    for e in entries:
        if e["sheet"] not in (FR, DORM):
            continue
        res = residence(e["building"]) or residence(e["room"])
        d1 = alnum(e["name"])
        d2 = alnum(e["name"] + (e["model"].split(" · ", 1)[1] if " · " in e["model"] else ""))
        cands = [m for m in master_furn if same_residence(res, m["res"])
                 and (m["desc"] in (d1, d2) or (len(d1) >= 6 and m["desc"].startswith(d1)))]
        if cands:
            overlaps.append((e, cands))
    # House Equip itemises flat appliances that the master list also counts in bulk ("SR&FR" x69).
    appliance_overlaps = []
    for label, rx in APPLIANCE_TYPES:
        rx = re.compile(rx, re.I)
        he = [e for e in entries if e["sheet"] == HOUSE and rx.search(e["name"])]
        ms = [e for e in entries if e["sheet"] == MASTER and e["category"] == "Equipment" and rx.search(e["name"])
              and (residence(e["building"]) or re.search(r"\b(D?FR|SR)", e["building"].upper()))]
        if he and ms:
            appliance_overlaps.append((label, he, ms))

    # --- output: register.js
    def asset(e):
        a = {"id": e["id"], "name": e["name"], "system": e["system"], "category": e["category"],
             "building": e["building"], "location": e["location"], "room": e["room"], "model": e["model"],
             "serial": e["serial"], "tag": e["tag"], "qty": e["units"], "year": e.get("year"),
             "life": e.get("life"), "lifePct": e.get("lifePct"), "criticality": e.get("criticality"),
             "amc": e.get("amc"), "ppm": e.get("ppm"), "freq": e["freq"], "src": src(e), "note": " ".join(e.get("notes", []))}
        # A status word typed into the serial column is shown as a note, not as a serial number.
        if a["serial"].upper() in ("DAMAGED", "SHIFTED"):
            a["note"] = ("Register marks this item: " + a["serial"].capitalize() + ". " + a["note"]).strip()
            a["serial"] = ""
        return {k: v for k, v in a.items() if v not in (None, "")}
    assets = [asset(e) for e in entries]
    counts = {"rows": len(records), "assets": len(assets), "removedDuplicates": len(removed),
              "mergedRows": sum(len(e["rows"]) - 1 for e in merged)}
    head = {"source": SRC.name, "title": "Master Asset Register · FM · Sep 2026",
            "generated": datetime.date.today().isoformat(), "counts": counts}
    dump = lambda o: json.dumps(o, ensure_ascii=False, separators=(",", ":"))
    js = ("// Generated by tools/import_register.py from data/" + SRC.name + ". Do not edit by hand.\n"
          "window.SOP_REGISTER = " + dump(head)[:-1] + ',"assets":[\n' +
          ",\n".join(dump(a) for a in assets) + "\n]};\n")

    # --- output: report
    rep = openpyxl.Workbook()
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0D7A4C")

    def sheet(title, header, rows, widths):
        ws = rep.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
        return ws

    by_sheet_assets = collections.Counter(e["sheet"] for e in entries)
    by_prefix = collections.Counter(e["id"].split("-")[0] for e in entries)
    removed_units = sum(r["rec"]["units"] for r in removed)
    final_units = sum(e["units"] for e in entries)
    summary = [["Source workbook", SRC.name], ["Generated", head["generated"]], [],
               ["Sheet", "Rows read", "Units (Quantity, blank/0 = 1)", "Duplicates removed", "Rows merged away",
                "Assets output", "Units output"]]
    for s in SHEETS:
        summary.append([s, rows_read[s], sum(r["units0"] for r in records if r["sheet"] == s),
                        sum(1 for r in removed if r["rec"]["sheet"] == s),
                        sum(len(e["rows"]) - 1 for e in merged if e["sheet"] == s),
                        by_sheet_assets[s], sum(e["units"] for e in entries if e["sheet"] == s)])
    summary += [["Total", len(records), total_units, len(removed), counts["mergedRows"], len(entries), final_units],
                [EXCLUDED, "excluded", "Residents' names/addresses (personal data), not assets"], [],
                ["Check", f"{final_units} units output + {removed_units} removed duplicate units = "
                          f"{final_units + removed_units} (source total {total_units})"],
                ["Kept with note", sum(1 for e in entries if e.get("notes")),
                 f"assets that share a serial/tag with another asset ({len(collisions)} pairs): kept, see "
                 "'Suspected entry errors (kept)'"],
                ["Merged entries", len(merged), "entries built from 2+ identical rows"],
                ["Possible overlaps", len(overlaps), "FR Furniture / dorm parents entries with look-alike master furniture "
                 "in the same residence (kept both)"], [],
                ["ID prefix", "Assets"]] + [[p, by_prefix[p]] for p in sorted(by_prefix)] + \
               [[], ["System", "Assets", "Units"]] + \
               [[s, n, sum(e["units"] for e in entries if e["system"] == s)]
                for s, n in collections.Counter(e["system"] for e in entries).most_common()] + \
               [[], ["Category", "Assets", "Units"]] + \
               [[c, n, sum(e["units"] for e in entries if e["category"] == c)]
                for c, n in collections.Counter(e["category"] for e in entries).most_common()]
    ws = sheet("Summary", ["Asset register import — duplicate check", ""], summary, [26, 14, 30, 20, 18, 14, 14])
    rep.remove(rep["Sheet"])
    for r in (5, 5 + len(SHEETS) + 1):
        for c in ws[r]:
            c.font = Font(bold=True)
    sheet("Duplicates removed",
          ["Removed row", "Description", "Place", "Serial", "Tag", "Qty", "Kept asset id", "Kept row",
           "Kept description", "Kept place", "Matched on", "Note"],
          [[src(r["rec"]), r["rec"]["name"], place(r["rec"]), r["rec"]["serial"], r["rec"]["tag"], r["rec"]["units"],
            r["kept"]["id"], src(r["kept"]) if len(r["kept"]["rows"]) > 1 else f"{r['kept']['sheet']} r{r['kept']['rows'][0]}",
            r["kept"]["name"], place(r["kept"]), r["match"], r["note"]] for r in removed],
          [24, 26, 30, 18, 16, 6, 12, 22, 26, 30, 26, 70])
    sheet("Suspected entry errors (kept)",
          ["Asset id", "Source", "Description", "Place", "Other asset id", "Other source", "Other description",
           "Other place", "Shared value", "Why judged separate units", "Action"],
          [[a["id"], src(a), a["name"], place(a), b["id"], src(b), b["name"], place(b),
            ", ".join(f"{k} {v}" for k, v in sh), "; ".join(why),
            "Kept both with a note on each asset; verify the serial/tag on site"] for a, b, sh, why in collisions],
          [10, 24, 24, 30, 10, 24, 24, 30, 26, 60, 34])
    sheet("Merged units", ["Asset id", "Description", "Building", "Location", "Room", "Model", "Tag", "Qty",
                           "Rows merged", "Source rows"],
          [[e["id"], e["name"], e["building"], e["location"], e["room"], e["model"], e["tag"], e["units"],
            len(e["rows"]), src(e)] for e in merged],
          [12, 40, 22, 26, 12, 20, 26, 6, 8, 40])
    sheet("Possible overlaps", ["Asset id", "Source", "Building", "Room", "Area", "Description", "Qty",
                                "Master look-alikes (ids)", "Master units", "Master places", "Action"],
          [[e["id"], src(e), e["building"], e["room"], e["location"], e["name"], e["units"],
            ", ".join(m["id"] for m in c[:12]) + (f" … +{len(c) - 12}" if len(c) > 12 else ""),
            sum(m["units"] for m in c), "; ".join(sorted({place(m) for m in c}))[:300],
            "Kept both: no shared serial/tag, match is by residence and description only"]
           for e, c in overlaps],
          [10, 22, 18, 14, 16, 30, 6, 40, 10, 50, 40])
    ws = rep["Possible overlaps"]
    ws.append([])
    ws.append(["House Equip vs Asset Master List: residence appliances (no shared serial/tag, so nothing removed)"])
    ws.append(["Type", "House Equip assets", "House Equip units", "House Equip ids", "Master assets", "Master units",
               "Master ids", "Master places (sample)", "Action"])
    for c in ws[ws.max_row - 1] + ws[ws.max_row]:
        c.font = Font(bold=True)
    for label, he, ms in appliance_overlaps:
        ws.append([label, len(he), sum(e["units"] for e in he), f"{he[0]['id']} … {he[-1]['id']}", len(ms),
                   sum(e["units"] for e in ms), ", ".join(e["id"] for e in ms[:8]) + (" …" if len(ms) > 8 else ""),
                   "; ".join(sorted({place(e) for e in ms}))[:200],
                   "Kept both: check whether the master entries are the same units as House Equip"])

    # --- privacy guard: no occupant / resident name may reach an output
    rep_text = "\n".join(str(c.value) for w in rep.worksheets for row in w.iter_rows() for c in row if c.value)
    for name in private_names:
        rx = re.compile(r"(?<![A-Za-z])" + re.escape(name) + r"(?![A-Za-z])", re.I)
        assert len(name) < 3 or not (rx.search(js) or rx.search(rep_text)), "personal name leaked into output"

    OUT_JS.write_text(js, encoding="utf-8")
    OUT_REPORT.parent.mkdir(exist_ok=True)
    rep.save(OUT_REPORT)

    # --- console summary
    print(f"Read {len(records)} rows ({', '.join(f'{s}: {rows_read[s]}' for s in SHEETS)}); skipped {EXCLUDED}")
    print(f"Kept {len(collisions)} serial/tag collisions as separate assets "
          f"({sum(1 for e in entries if e.get('notes'))} assets carry a note)")
    print(f"Removed {len(removed)} duplicates ({removed_units} units); merged {counts['mergedRows']} rows "
          f"into {len(merged)} entries; {len(overlaps)} possible furniture overlaps kept")
    for label, he, ms in appliance_overlaps:
        print(f"  overlap? {label}: House Equip {sum(e['units'] for e in he)} units vs master "
              f"{sum(e['units'] for e in ms)} units in residences")
    print(f"Units: {final_units} output + {removed_units} removed = {final_units + removed_units} "
          f"(source {total_units})")
    print("Systems:", dict(collections.Counter(e["system"] for e in entries).most_common()))
    general = collections.Counter(e["name"] for e in entries if e["system"] == "General")
    if general:
        print("Top General:", general.most_common(15))
    print(f"Wrote {OUT_JS.relative_to(ROOT)} ({OUT_JS.stat().st_size // 1024} KB, {len(assets)} assets) and "
          f"{OUT_REPORT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
